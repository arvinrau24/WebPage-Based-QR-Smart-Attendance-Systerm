import openpyxl
from io import BytesIO
import os

def clean(val):
    return str(val).strip() if val is not None else ''

def parse_student_list(file_bytes, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'):
    """Parse student list from Excel, PDF, Word or Image."""

    if 'pdf' in content_type:
        return parse_from_pdf(file_bytes)
    elif 'word' in content_type or 'docx' in content_type or content_type in [
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'application/msword'
    ]:
        return parse_from_docx(file_bytes)
    elif 'image' in content_type:
        return parse_from_image(file_bytes, content_type)
    else:
        return parse_from_excel(file_bytes)


def parse_from_excel(file_bytes):
    wb = openpyxl.load_workbook(BytesIO(file_bytes))
    results = []

    for sheet in wb.worksheets:
        course_code = None
        section = None
        students = []
        reading_students = False

        for row in sheet.iter_rows(values_only=True):
            cells = list(row)

            if len(cells) < 3:
                continue

            if clean(cells[1]) == 'KURSUS' and len(cells) > 3 and cells[3]:
                val = clean(cells[3]).lstrip(':').strip()
                course_code = val.split('-')[0].strip()

            if 'SEKSYEN' in clean(cells[1]).upper() and len(cells) > 3 and cells[3]:
                val = clean(cells[3]).lstrip(':').strip()
                section = val

            if clean(cells[1]) == 'BIL.' and 'MATRIK' in clean(cells[2]).upper():
                reading_students = True
                continue

            if reading_students:
                matric = clean(cells[2])
                name = clean(cells[3])
                phone = clean(cells[5]) if len(cells) > 5 else ''
                email = clean(cells[6]) if len(cells) > 6 else ''

                if matric and matric[0] in ['B', 'A', 'b', 'a'] and name:
                    students.append({
                        'matric_number': matric.upper(),
                        'full_name': name,
                        'phone': phone,
                        'email': email if email else f"{matric.lower()}@student.utem.edu.my"
                    })

        if course_code and students:
            results.append({
                'course_code': course_code,
                'section': section or 'N/A',
                'students': students
            })

    return results


def parse_from_pdf(file_bytes):
    import pdfplumber
    results = []

    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ''
            students = extract_students_from_text(text)
            course_code, section = extract_course_info_from_text(text)
            if course_code and students:
                results.append({'course_code': course_code, 'section': section, 'students': students})

    return results


def parse_from_docx(file_bytes):
    from docx import Document
    doc = Document(BytesIO(file_bytes))
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    students = extract_students_from_text(full_text)
    course_code, section = extract_course_info_from_text(full_text)

    if course_code and students:
        return [{'course_code': course_code, 'section': section, 'students': students}]
    return []


def parse_from_image(file_bytes, content_type):
    """Use Groq vision to read student list from image."""
    import base64
    import json
    import re
    from groq import Groq

    client = Groq(api_key=os.getenv('GROQ_API_KEY'))
    image_base64 = base64.b64encode(file_bytes).decode('utf-8')
    ext = content_type.split('/')[-1]

    prompt = """
This is a university student list. Extract all students.
Return ONLY a JSON object like:
{
  "course_code": "BERR4713",
  "section": "1/1",
  "students": [
    {"matric_number": "B122320018", "full_name": "ARVIN RAU", "phone": "0123456789", "email": "b122320018@student.utem.edu.my"}
  ]
}
No extra text, no markdown, no backticks.
"""

    response = client.chat.completions.create(
        model="meta-llama/llama-4-scout-17b-16e-instruct",
        messages=[{
            "role": "user",
            "content": [
                {"type": "image_url", "image_url": {"url": f"data:image/{ext};base64,{image_base64}"}},
                {"type": "text", "text": prompt}
            ]
        }],
        max_tokens=2000,
    )

    raw = response.choices[0].message.content.strip()
    match = re.search(r'\{.*\}', raw, re.DOTALL)
    if match:
        data = json.loads(match.group(0))
        return [data]
    return []


def extract_course_info_from_text(text):
    import re
    course_code = None
    section = None

    course_match = re.search(r'KURSUS\s*:?\s*([A-Z]{4}\d{4})', text)
    if course_match:
        course_code = course_match.group(1)

    section_match = re.search(r'SEKSYEN\s*[/\w]*\s*:?\s*(\d+/\d+|\d+)', text)
    if section_match:
        section = section_match.group(1)

    return course_code, section or 'N/A'


def extract_students_from_text(text):
    import re
    students = []
    lines = text.split('\n')

    for line in lines:
        matric_match = re.search(r'\b([BA]\d{9,})\b', line)
        if matric_match:
            matric = matric_match.group(1).upper()
            remaining = line[matric_match.end():].strip()
            parts = remaining.split()
            name = ' '.join([p for p in parts if not p.startswith('01') and '@' not in p and not p.isdigit()][:4])
            phone = next((p for p in parts if p.startswith('01')), '')
            email = next((p for p in parts if '@' in p), f"{matric.lower()}@student.utem.edu.my")

            if name:
                students.append({
                    'matric_number': matric,
                    'full_name': name,
                    'phone': phone,
                    'email': email
                })

    return students